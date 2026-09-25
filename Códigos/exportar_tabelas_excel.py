"""
exportar_tabelas_excel.py  (passo 16 do pipeline — SEMPRE rodar apos recalculos)
------------------------------------------------------------------
TCC MBA USP/Esalq - Credito Verde e IA para PMEs
Autor: Helio Vinicius Moreira Ribeiro

Consolida TODAS as tabelas do Resultado Preliminar num unico Excel editavel,
com GRAFICOS NATIVOS do Excel (editaveis pelo autor), lendo as saidas vigentes
do pipeline (nao hardcoda resultados; recalcula da base pontuada e dos CSVs).

Regra do projeto: dados e graficos de resultado devem SEMPRE sair tambem em
Excel — este script e o passo que garante isso. Rode-o ao final de qualquer
recalculo:  python -B Códigos/exportar_tabelas_excel.py

OUTPUT:
  "20260612 - Tabelas_e_Graficos_TCC - claude.xlsx"  (RAIZ, para edicao pelo autor)
  Convencao de nomes dos documentos: aaaammdd - nome - autor.ext (data de CRIACAO)
"""

import os
import re
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DADOS    = os.path.join(BASE_DIR, "Dados")
SCORED   = os.path.join(DADOS, "base_analitica_scored.csv")
SAIDA    = os.path.join(BASE_DIR, "20260612 - Tabelas_e_Graficos_TCC - claude.xlsx")

REGIAO = {
    **{u: "Norte"        for u in ["AC","AP","AM","PA","RO","RR","TO"]},
    **{u: "Nordeste"     for u in ["AL","BA","CE","MA","PB","PE","PI","RN","SE"]},
    **{u: "Centro-Oeste" for u in ["DF","GO","MT","MS"]},
    **{u: "Sudeste"      for u in ["ES","MG","RJ","SP"]},
    **{u: "Sul"          for u in ["PR","RS","SC"]},
}
ORD_REG   = ["Norte", "Nordeste", "Centro-Oeste", "Sudeste", "Sul"]
ORD_FAIXA = ["Verde-A", "Verde-B", "Amarelo", "Vermelho", "Vetado", "Fora de escopo"]
PORTE_NOME = {"01": "Micro", "03": "Pequena", "05": "Media+", "00": "ND"}

# estilos
F_TIT  = Font(name="Arial", size=12, bold=True, color="1A237E")
F_HDR  = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_TXT  = Font(name="Arial", size=10)
F_NOTA = Font(name="Arial", size=9, italic=True, color="555555")
FILL_H = PatternFill("solid", start_color="1A237E")
PCT, NUM0, NUM1, NUM3 = "0.0%", "#,##0", "0.0", "0.000"


def carregar_base():
    print("[1/4] Lendo base pontuada (colunas de tabela) ...")
    cols = ["uf", "porte", "faixa", "score_e", "score_s", "score_enquadramento",
            "score_socioambiental", "flag_economia_verde", "fbb_ev_eixo",
            "flag_risco_ambiental", "flag_clima", "qtd_infracoes", "embargado",
            "ceis_sancoes_ativas", "salario_medio_sm", "rotatividade",
            "gap_genero_ajustado_cnae", "gap_raca_ajustado_cnae",
            "pct_chefia_feminina_cnae", "dispersao_salarial_cnae"]
    df = pd.read_csv(SCORED, sep=";", dtype=str, encoding="utf-8-sig",
                     usecols=lambda c: c in cols, low_memory=False)
    num = [c for c in cols if c not in ("uf", "porte", "faixa", "fbb_ev_eixo")]
    for c in num:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df["regiao"] = df["uf"].astype(str).str.strip().str.upper().map(REGIAO).fillna("ND")
    df["porte_nome"] = df["porte"].astype(str).str.strip().map(PORTE_NOME).fillna("ND")
    df["aprovado"] = df["faixa"].isin(["Verde-A", "Verde-B"]).astype(int)
    return df


def parse_relatorio_classificacao():
    """Extrai a tabela de desempenho da Fase 2 do relatorio vigente.
    Robusto aos dois formatos: 4 metricas (sem F1) ou 5 (com F1)."""
    cam = os.path.join(DADOS, "relatorio_classificacao.txt")
    linhas = []
    if os.path.exists(cam):
        for ln in open(cam, encoding="utf-8"):
            m = re.match(r"\s{2}(Logistica|XGBoost)((?:\s+[\d.]+)+)\s*$", ln)
            if m:
                nums = [float(x) for x in m.group(2).split()]
                linhas.append([m.group(1)] + nums)
    if linhas and len(linhas[0]) == 6:
        cols = ["Modelo", "Acuracia", "Precisao", "Recall", "F1", "AUC-ROC"]
    else:
        cols = ["Modelo", "Acuracia", "Precisao", "Recall", "AUC-ROC"]
    return pd.DataFrame(linhas, columns=cols)


def parse_roc_binaria():
    cam = os.path.join(DADOS, "relatorio_roc_binaria.txt")
    out = {}
    pats = {"AUC-ROC (binaria)": r"AUC-ROC\s*:\s*([\d.]+)",
            "Ponto de corte (Youden)": r"corte \(Youden\):\s*([\d.]+)",
            "Sensibilidade no corte": r"Sensibilidade \(recall\)\s*:\s*([\d.]+)",
            "Especificidade no corte": r"Especificidade\s*:\s*([\d.]+)",
            "Precisao no corte": r"Precisao\s*:\s*([\d.]+)",
            "Acuracia no corte": r"Acuracia\s*:\s*([\d.]+)"}
    if os.path.exists(cam):
        txt = open(cam, encoding="utf-8").read()
        for k, p in pats.items():
            m = re.search(p, txt)
            if m:
                out[k] = float(m.group(1))
    return out


def escrever_aba(ws, titulo, df, formatos=None, nota=None, largura0=24):
    """Escreve titulo + DataFrame formatado; retorna linha do header (p/ graficos)."""
    ws["A1"] = titulo; ws["A1"].font = F_TIT
    hdr_row = 3
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=hdr_row, column=j, value=str(col))
        c.font = F_HDR; c.fill = FILL_H
        c.alignment = Alignment(horizontal="center")
    for i, (_, row) in enumerate(df.iterrows(), start=hdr_row + 1):
        for j, col in enumerate(df.columns, start=1):
            v = row[col]
            c = ws.cell(row=i, column=j,
                        value=(None if pd.isna(v) else
                               v.item() if hasattr(v, "item") else v))
            c.font = F_TXT
            if formatos and col in formatos:
                c.number_format = formatos[col]
    ws.column_dimensions["A"].width = largura0
    for j in range(2, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16
    if nota:
        c = ws.cell(row=hdr_row + len(df) + 2, column=1, value=nota)
        c.font = F_NOTA
    return hdr_row


def grafico_barras(ws, titulo, hdr_row, n_rows, col_dados, ancora, n_series=1, pct=False):
    ch = BarChart(); ch.type = "col"; ch.title = titulo; ch.style = 10
    ch.height = 8; ch.width = 16; ch.gapWidth = 60
    data = Reference(ws, min_col=col_dados, max_col=col_dados + n_series - 1,
                     min_row=hdr_row, max_row=hdr_row + n_rows)
    cats = Reference(ws, min_col=1, min_row=hdr_row + 1, max_row=hdr_row + n_rows)
    ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
    if pct:
        ch.y_axis.numFmt = "0%"
    ws.add_chart(ch, ancora)


def main():
    print("=" * 64)
    print(" EXPORTACAO DE TABELAS E GRAFICOS PARA EXCEL (passo 16)")
    print("=" * 64)
    df = carregar_base()
    n = len(df)
    wb = Workbook(); wb.remove(wb.active)

    print("[2/4] Montando tabelas ...")

    # ---- Leia-me ----
    ws = wb.create_sheet("Leia-me")
    ws["A1"] = "Tabelas e Graficos do Resultado Preliminar — Bureau de Credito Verde"; ws["A1"].font = F_TIT
    infos = [
        "Gerado por Códigos/exportar_tabelas_excel.py a partir das saidas VIGENTES do pipeline.",
        "Re-rodar o script apos qualquer recalculo: python -B Códigos/exportar_tabelas_excel.py",
        "Cada aba = uma tabela do esboco (numeracao identica). Graficos sao nativos do Excel (editaveis).",
        "Valores numericos reais (nao texto): formate/edite a vontade para o documento final.",
        f"Base: base_analitica_scored.csv ({n:,} registros).",
    ]
    for i, t in enumerate(infos, start=3):
        ws.cell(row=i, column=1, value=t).font = F_TXT
    ws.column_dimensions["A"].width = 110

    # ---- T1 Universo ----
    t1 = pd.DataFrame({
        "Indicador": ["Empresas na amostra (AAS 10% RFB, ativas)",
                      "Porte: Micro", "Porte: Pequena", "Porte: Media+",
                      "Com autuacao IBAMA", "Com embargo IBAMA ativo", "Com sancao CEIS ativa"],
        "Valor": [n,
                  int((df.porte_nome == "Micro").sum()), int((df.porte_nome == "Pequena").sum()),
                  int((df.porte_nome == "Media+").sum()),
                  int((df.qtd_infracoes.fillna(0) > 0).sum()),
                  int((df.embargado.fillna(0) > 0).sum()),
                  int((df.ceis_sancoes_ativas.fillna(0) > 0).sum())],
        "% da amostra": [1.0,
                         (df.porte_nome == "Micro").mean(), (df.porte_nome == "Pequena").mean(),
                         (df.porte_nome == "Media+").mean(),
                         (df.qtd_infracoes.fillna(0) > 0).mean(),
                         (df.embargado.fillna(0) > 0).mean(),
                         (df.ceis_sancoes_ativas.fillna(0) > 0).mean()],
    })
    ws = wb.create_sheet("T1_Universo")
    escrever_aba(ws, "Tabela 1. Universo analisado e ocorrencias socioambientais", t1,
                 {"Valor": NUM0, "% da amostra": PCT}, "Fonte: resultados originais da pesquisa", 42)

    # ---- T2 FEBRABAN ----
    fe = df.flag_economia_verde.fillna(0) > 0
    t2 = pd.DataFrame({
        "Marcacao": ["Economia verde (qualquer eixo)", "  eixo Social", "  eixo Ambiental",
                     "  eixos Social + Ambiental", "Exposicao ao risco ambiental",
                     "Exposicao as mudancas climaticas"],
        "Empresas": [int(fe.sum()),
                     int((df.fbb_ev_eixo == "Social").sum()),
                     int((df.fbb_ev_eixo == "Ambiental").sum()),
                     int((df.fbb_ev_eixo == "Social + Ambiental").sum()),
                     int((df.flag_risco_ambiental.fillna(0) > 0).sum()),
                     int((df.flag_clima.fillna(0) > 0).sum())],
    })
    t2["%"] = t2["Empresas"] / n
    ws = wb.create_sheet("T2_FEBRABAN")
    h = escrever_aba(ws, "Tabela 2. Atividade verde e exposicoes — Taxonomia FEBRABAN", t2,
                     {"Empresas": NUM0, "%": PCT},
                     "Fonte: Taxonomia Verde FEBRABAN (Dez/2020); casamento por subclasse CNAE = 100%", 34)
    grafico_barras(ws, "Atividade verde e exposicoes (empresas)", h, len(t2), 2, "E3")

    # ---- T3 Equidade ----
    eq_cols = {"salario_medio_sm": ("Salario medio setorial (SM)", NUM1),
               "rotatividade": ("Rotatividade setorial", NUM3),
               "gap_genero_ajustado_cnae": ("Gap salarial de genero ajustado", NUM3),
               "gap_raca_ajustado_cnae": ("Gap salarial de raca ajustado", NUM3),
               "pct_chefia_feminina_cnae": ("% de mulheres em cargos de chefia", PCT),
               "dispersao_salarial_cnae": ("Dispersao salarial (CV)", NUM3)}
    t3 = pd.DataFrame({"Indicador": [v[0] for v in eq_cols.values()],
                       "Media": [df[k].mean() for k in eq_cols]})
    fmt3 = {"Media": NUM3}
    ws = wb.create_sheet("T3_Equidade")
    escrever_aba(ws, "Tabela 3. Indicadores setoriais de equidade (medias por CNAE)", t3, fmt3,
                 "Fonte: RAIS Vinculos (MTE); gaps ajustados por escolaridade e tempo de emprego", 38)

    # ---- T4 Faixas ----
    vc = df.faixa.value_counts()
    t4 = pd.DataFrame({"Faixa": ORD_FAIXA,
                       "Empresas": [int(vc.get(f, 0)) for f in ORD_FAIXA]})
    t4["%"] = t4["Empresas"] / n
    ws = wb.create_sheet("T4_Faixas")
    h = escrever_aba(ws, "Tabela 4. Distribuicao por faixa de enquadramento", t4,
                     {"Empresas": NUM0, "%": PCT}, "Fonte: resultados originais da pesquisa")
    grafico_barras(ws, "Empresas por faixa de enquadramento", h, len(t4), 2, "E3")

    # ---- T5 Score por porte ----
    g = df.groupby("porte_nome")[["score_e", "score_s", "score_enquadramento"]].mean()
    g = g.reindex(["Media+", "Pequena", "Micro"]).dropna(how="all").reset_index()
    g.columns = ["Porte", "Pilar E", "Pilar S", "Total"]
    ws = wb.create_sheet("T5_Score_Porte")
    h = escrever_aba(ws, "Tabela 5. Score medio da rubrica por porte (0-100)", g,
                     {"Pilar E": NUM1, "Pilar S": NUM1, "Total": NUM1},
                     "Fonte: resultados originais da pesquisa")
    grafico_barras(ws, "Score medio por porte (E, S, Total)", h, len(g), 2, "F3", n_series=3)

    # ---- T6 Fase 2 ----
    t6 = parse_relatorio_classificacao()
    ws = wb.create_sheet("T6_Fase2")
    h = escrever_aba(ws, "Tabela 6. Desempenho da Fase 2 (teste; 4 classes)", t6,
                     {c: NUM3 for c in t6.columns if c != "Modelo"},
                     "Fonte: relatorio_classificacao.txt (SMOTE so no treino; teste intocado)")
    if len(t6):
        grafico_barras(ws, "Fase 2 — metricas por modelo", h, len(t6), 2, "H3",
                       n_series=len(t6.columns) - 1)

    # ---- T7 ROC binaria + score ----
    roc = parse_roc_binaria()
    sc = df.loc[df.score_socioambiental >= 0, "score_socioambiental"]
    t7 = pd.DataFrame({"Indicador": list(roc.keys()) +
                       ["Score socioambiental medio (0-1000)", "Aprovados (Verde-A/B)", "% aprovados"],
                       "Valor": list(roc.values()) +
                       [sc.mean(), int(df.aprovado.sum()), df.aprovado.mean()]})
    ws = wb.create_sheet("T7_ROC_Score")
    escrever_aba(ws, "Tabela 7. Decisao binaria de elegibilidade e score", t7, {"Valor": "0.000"},
                 "Fonte: relatorio_roc_binaria.txt + base pontuada (formate '% aprovados' como % se preferir)", 38)

    # ---- T8 Apendice ----
    cam8 = os.path.join(DADOS, "apendice_comparacao_modelos.csv")
    if os.path.exists(cam8):
        t8 = pd.read_csv(cam8, sep=";")
        keep = [c for c in ["Modelo", "AUC", "F1_macro", "Acuracia"] if c in t8.columns]
        t8 = t8[keep]
        ws = wb.create_sheet("T8_Apendice")
        h = escrever_aba(ws, "Tabela 8. Comparacao multi-modelo (apendice)", t8,
                         {c: NUM3 for c in keep if c != "Modelo"},
                         "Fonte: apendice_comparacao_modelos.csv (Optuna; SHAP)")
        grafico_barras(ws, "Apendice — AUC por modelo", h, len(t8), 2, "F3")

    # ---- T9 Regional ----
    def agg_reg(gg):
        return pd.Series({"Empresas": len(gg), "% amostra": len(gg) / n,
                          "Score medio": gg.score_enquadramento.mean(),
                          "% aprovados": gg.aprovado.mean()})
    t9 = df.groupby("regiao").apply(agg_reg, include_groups=False).reindex(ORD_REG + ["ND"]).dropna(how="all")
    t9 = t9.reset_index().rename(columns={"regiao": "Regiao"})
    t9["Empresas"] = t9["Empresas"].astype(int)
    ws = wb.create_sheet("T9_Regional")
    h = escrever_aba(ws, "Tabela 9. Amostra e rubrica por regiao", t9,
                     {"Empresas": NUM0, "% amostra": PCT, "Score medio": NUM1, "% aprovados": PCT},
                     "Fonte: resultados originais da pesquisa")
    grafico_barras(ws, "% aprovados (Verde-A/B) por regiao", h, len(t9), 5, "G3", pct=True)

    # ---- T10 Fairness ----
    cam10 = os.path.join(DADOS, "tabela_regional_fairness.csv")
    if os.path.exists(cam10):
        t10 = pd.read_csv(cam10, sep=";", encoding="utf-8-sig")
        ws = wb.create_sheet("T10_Fairness")
        fmt = {c: NUM3 for c in t10.columns if c not in ("regiao", "n_teste")}
        fmt["n_teste"] = NUM0
        h = escrever_aba(ws, "Tabela 10. Fase 2 por regiao (teste) — igualdade de oportunidade",
                         t10, fmt, "Fonte: analise_regional.py (mesmo split de teste da Fase 2)")
        col_rec = list(t10.columns).index("recall (sensibilidade)") + 1
        grafico_barras(ws, "Recall e FPR por regiao (Fase 2, teste)", h, len(t10), col_rec, "I3", n_series=2)

    # ---- UFs (apoio) ----
    cam_uf = os.path.join(DADOS, "tabela_regional_uf.csv")
    if os.path.exists(cam_uf):
        tuf = pd.read_csv(cam_uf, sep=";", encoding="utf-8-sig")
        ws = wb.create_sheet("Apoio_UF")
        fmt = {c: NUM3 for c in tuf.columns if c not in ("regiao", "uf", "empresas")}
        fmt["empresas"] = NUM0
        escrever_aba(ws, "Apoio. Distribuicao e score por UF", tuf, fmt,
                     "Fonte: analise_regional.py — base para mapas/recortes que o autor quiser")

    # ---- T11 Desenho amostral: aderencia populacao x amostra ----
    cam_ader = os.path.join(DADOS, "tabela_aderencia_amostral.csv")
    if os.path.exists(cam_ader):
        tad = pd.read_csv(cam_ader, sep=";", encoding="utf-8-sig")
        ws = wb.create_sheet("T11_Aderencia_Amostral")
        fmt = {}
        for c in tad.columns:
            if c in ("populacao", "amostra"):
                fmt[c] = NUM0
            elif c.endswith("_%"):
                fmt[c] = PCT
            elif c in ("dif_pp", "fracao_%", "erro_amostral_pp"):
                fmt[c] = NUM3
        escrever_aba(ws, "Tabela 11. Aderencia da amostra a populacao (por estrato)", tad, fmt,
                     "Fonte: analise_amostral.py — varredura da base bruta da RFB; "
                     "fracao_% proxima de 10 indica proporcionalidade", 26)

    # ---- T12 Estimativas populacionais (peso amostral) ----
    cam_est = os.path.join(DADOS, "estimativas_populacionais.csv")
    if os.path.exists(cam_est):
        tes = pd.read_csv(cam_est, sep=";", encoding="utf-8-sig")
        mostrar = [c for c in ["dominio", "indicador", "amostra", "p", "total",
                               "tot_li", "tot_ls", "erro_pp"] if c in tes.columns]
        tes = tes[mostrar].rename(columns={
            "p": "proporcao", "total": "estimativa_populacional",
            "tot_li": "IC95_inferior", "tot_ls": "IC95_superior",
            "erro_pp": "margem_erro_pp"})
        ws = wb.create_sheet("T12_Estim_Populacionais")
        fmt = {"amostra": NUM0, "proporcao": PCT, "estimativa_populacional": NUM0,
               "IC95_inferior": NUM0, "IC95_superior": NUM0, "margem_erro_pp": NUM3}
        h = escrever_aba(ws, "Tabela 12. Estimativas para a populacao (amostra expandida pelo peso)",
                         tes, fmt,
                         "Fonte: estimativas_populacionais.py — AAS sem reposicao, "
                         "correcao (1-f) para populacao finita, IC 95%", 24)
        brasil = tes[tes["dominio"] == "Brasil"] if "dominio" in tes.columns else tes
        if len(brasil):
            col_est = mostrar.index("total") + 1 if "total" in mostrar else 5
            grafico_barras(ws, "Estimativa populacional por indicador (Brasil)",
                           h, len(brasil), col_est, "J3")

    print("[3/4] Salvando workbook ...")
    wb.save(SAIDA)
    print(f"[4/4] Salvo: {SAIDA}")
    print("\nLembrete: rode este script apos QUALQUER recalculo do pipeline.")


if __name__ == "__main__":
    main()
