"""
construir_base_febraban.py  —  ARQUIVO ÚNICO, fonte única.

Lê a Taxonomia Verde da FEBRABAN DIRETO do xlsx oficial (não usa CSV intermediário),
normaliza tudo em código (inclusive a grafia [Social+Ambiental] == [Social + Ambiental]),
casa na base analítica por subclasse CNAE e grava a base enriquecida.

Modelo de duas camadas:
  ATIVIDADE (FEBRABAN, por CNAE)  ->  fbb_economia_verde / fbb_ev_eixo / fbb_risco_ambiental / fbb_clima
  CONDUTA   (seus dados, por firma) ->  IBAMA, social RAIS/CAGED, veto CEIS, % cadeia verde  [NÃO tocado aqui]
Escopo: crédito sustentável/ESG — 'economia verde' inclui o eixo Social (linha FEBRABAN).
Recorte: versão Dez/2020, vigente em 31/05/2026 (monitorar nova edição 2026).

A proteção do xlsx é de EDIÇÃO, não de leitura: openpyxl lê normalmente. Não desproteja o arquivo.
"""
import re
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

# ===================== CAMINHOS ROBUSTOS =====================
# Derivados da raiz do projeto (pasta-pai de Códigos/), não do diretório de
# trabalho — evita FileNotFoundError quando o script é chamado de outro CWD.
RAIZ          = Path(__file__).resolve().parents[1]
FEBRABAN_XLSX = str(RAIZ / "referencias" / "taxonomia_febraban" /
                    "FEBRABAN_Taxonomia_Lista_CNAEs_20210217.xlsx")
SHEET         = "Categorização_CNAE"
BASE_IN       = str(RAIZ / "Dados" / "base_analitica.csv")
BASE_OUT      = str(RAIZ / "Dados" / "base_analitica_febraban.csv")
CNAE_COL      = "cnae_fiscal_principal"   # subclasse de 7 dígitos da RFB (ex.: 0111301)
# CSVs do projeto: separador ';' e encoding utf-8-sig (mantém consistência a jusante).
CSV_SEP       = ";"
CSV_ENC       = "utf-8-sig"

# Colunas da aba (0-based), versão Dez/2020:
C_SUBCLASSE, C_DENOM, C_CLIMA, C_ECONVERDE, C_RISCO = 4, 5, 6, 10, 18
# =============================================================


def norm_cnae(x):
    """Qualquer formato de CNAE -> chave de 7 dígitos ('0111-3/01' -> '0111301')."""
    if pd.isna(x):
        return None
    d = re.sub(r"\D", "", str(x))
    return d.zfill(7) if d else None


def norm_ev(label):
    """Canoniza o rótulo de Economia Verde; colapsa [Social+Ambiental] e [Social + Ambiental]."""
    if not isinstance(label, str) or not label.strip():
        return ""
    s = re.sub(r"\s+", " ", label).strip()
    m = re.search(r"\[(.*?)\]", s)
    if not m:
        return s
    inner = m.group(1).lower()
    eixo = ("Social + Ambiental" if ("social" in inner and "ambiental" in inner)
            else "Social" if "social" in inner
            else "Ambiental" if "ambiental" in inner else "")
    nivel = ("Alta" if s.lower().startswith("alta")
             else "Moderada" if s.lower().startswith("moderada") else "")
    return f"{nivel} contribuição [{eixo}]" if (nivel and eixo) else s


def clean(x):
    return re.sub(r"\s+", " ", str(x)).strip() if x is not None else ""


def build_lookup(xlsx_path, sheet):
    """Extrai e normaliza a taxonomia direto do xlsx -> DataFrame por subclasse."""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[sheet]
    recs = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if len(r) <= C_SUBCLASSE or not r[C_SUBCLASSE]:
            continue
        key = re.sub(r"\D", "", clean(r[C_SUBCLASSE]))
        if len(key) != 7:          # só subclasse de 7 dígitos
            continue
        ev = norm_ev(clean(r[C_ECONVERDE]) if len(r) > C_ECONVERDE else "")
        eixo = ""
        mm = re.search(r"\[(.*?)\]", ev)
        if mm:
            eixo = mm.group(1)
        risco = clean(r[C_RISCO]) if len(r) > C_RISCO else ""
        clima = clean(r[C_CLIMA]) if len(r) > C_CLIMA else ""
        recs.append({
            "_cnae_key": key,
            "denominacao_febraban": clean(r[C_DENOM]),
            "fbb_economia_verde": ev,
            "fbb_ev_eixo": eixo,
            "fbb_risco_ambiental": risco,
            "fbb_clima": clima,
            "flag_economia_verde": 1 if ev else 0,
            "flag_risco_ambiental": 1 if risco else 0,
            "flag_clima": 1 if clima else 0,
        })
    look = pd.DataFrame(recs).drop_duplicates("_cnae_key")
    return look


def main():
    look = build_lookup(FEBRABAN_XLSX, SHEET)
    print(f"Taxonomia FEBRABAN carregada: {len(look)} subclasses")
    print("Economia verde por rótulo (normalizado):")
    print(look.loc[look.fbb_economia_verde != "", "fbb_economia_verde"]
          .value_counts().to_string())
    print()

    # dtype=str para TODAS as colunas: passthrough fiel da base (preserva zeros
    # a esquerda de 'porte' = '01'/'03'/'05', codigos CNAE, CNPJ etc.). A tipagem
    # numerica e feita a jusante (rubrica/Fase 2), nao aqui.
    base = pd.read_csv(BASE_IN, sep=CSV_SEP, encoding=CSV_ENC,
                       dtype=str, low_memory=False)
    base["_cnae_key"] = base[CNAE_COL].map(norm_cnae)

    fbb_cols = ["denominacao_febraban", "fbb_economia_verde", "fbb_ev_eixo",
                "fbb_risco_ambiental", "fbb_clima",
                "flag_economia_verde", "flag_risco_ambiental", "flag_clima"]
    m = base.merge(look, on="_cnae_key", how="left")
    for c in ["flag_economia_verde", "flag_risco_ambiental", "flag_clima"]:
        m[c] = m[c].fillna(0).astype(int)
    for c in ["denominacao_febraban", "fbb_economia_verde", "fbb_ev_eixo",
              "fbb_risco_ambiental", "fbb_clima"]:
        m[c] = m[c].fillna("")

    n = len(m)
    casadas = m["_cnae_key"].isin(look["_cnae_key"]).sum()
    print(f"Empresas: {n:,} | casaram com CNAE FEBRABAN: {casadas:,} ({casadas/n:.1%})")
    if casadas / n < 0.5:
        print("  *** baixa taxa: confira se", CNAE_COL, "é a subclasse de 7 dígitos ***")
    print(f"economia verde (qualquer eixo): {m['flag_economia_verde'].sum():,} ({m['flag_economia_verde'].mean():.1%})")
    print(f"  Social: {(m.fbb_ev_eixo=='Social').sum():,} | "
          f"Ambiental: {(m.fbb_ev_eixo=='Ambiental').sum():,} | "
          f"Social + Ambiental: {(m.fbb_ev_eixo=='Social + Ambiental').sum():,}")
    print(f"alto risco ambiental: {m['flag_risco_ambiental'].sum():,} | exposição climática: {m['flag_clima'].sum():,}")
    # Comparação verde ANTES (classificação ad-hoc por divisão CNAE) × AGORA (FEBRABAN).
    # A base não tem flag_verde; reconstruímos o "antes" pelas divisões do dicionário
    # antigo (stg_cnae_verde.csv, ou o fallback embutido de 6 divisões BNDES).
    VERDE_DIVISOES_ANTIGO = {"36", "37", "38", "39", "35", "02"}
    stg = RAIZ / "Dados" / "stg_cnae_verde.csv"
    if stg.exists():
        try:
            d = pd.read_csv(stg, sep=";", dtype={"cnae_divisao": str}, encoding="utf-8-sig")
            VERDE_DIVISOES_ANTIGO = {str(x).zfill(2) for x in d["cnae_divisao"]}
        except Exception:
            pass
    div = m[CNAE_COL].astype(str).str.replace(r"\D", "", regex=True).str.zfill(7).str[:2]
    antes = int(div.isin(VERDE_DIVISOES_ANTIGO).sum())
    agora = int(m["flag_economia_verde"].sum())
    print(f"verde ANTES (6 divisões BNDES {sorted(VERDE_DIVISOES_ANTIGO)}): {antes:,} "
          f"-> AGORA (FEBRABAN/ESG, subclasse): {agora:,}  ({agora-antes:+,})")

    m.drop(columns=["_cnae_key"]).to_csv(BASE_OUT, sep=CSV_SEP,
                                         index=False, encoding=CSV_ENC)
    print(f"\nGravado: {BASE_OUT}")


if __name__ == "__main__":
    main()
